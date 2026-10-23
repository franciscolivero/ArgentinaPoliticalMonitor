#!/usr/bin/env python3
"""
post_data.py — one place that turns the monitor's data files into (a) chart tables for
Datawrapper and (b) the numbers the Substack drafts quote. Used by publish_datawrapper.py and
draft_posts.py. Reads only files already in the repo.
"""

import csv, io, json, datetime, pathlib

HERE = pathlib.Path(__file__).parent
MONTHS_ES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
MONTHS_ES_LONG = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
                  "septiembre", "octubre", "noviembre", "diciembre"]


def load(name):
    p = HERE / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def table(header, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return buf.getvalue()


def at(dates, vals, target):
    """Last value on or before `target` (ISO date)."""
    best = None
    for d, v in zip(dates, vals):
        if d <= target and v is not None:
            best = v
    return best


def ago(days, ref=None):
    ref = ref or datetime.date.today()
    return (ref - datetime.timedelta(days=days)).isoformat()


def fdate(iso):
    d = datetime.date.fromisoformat(iso[:10])
    return f"{d.day} de {MONTHS_ES_LONG[d.month - 1]}"


def fmonth(iso):
    d = datetime.date.fromisoformat(iso[:10])
    return f"{MONTHS_ES_LONG[d.month - 1]} de {d.year}"


def pct(v, nd=0):
    return ("—" if v is None else f"{v:.{nd}f}".replace(".", ",")) + ("%" if v is not None else "")


def num(v, nd=2):
    return "—" if v is None else f"{v:.{nd}f}".replace(".", ",").replace("-", "−")


def delta(new, old, nd=1, unit=" pp"):
    if new is None or old is None:
        return "sin comparación"
    d = new - old
    if abs(d) < 0.5 * 10 ** -nd:
        return "sin cambios"
    return f"{'+' if d > 0 else '−'}{abs(d):.{nd}f}".replace(".", ",") + unit


# ---------------------------------------------------------------------------- Argentina tables
def _interp(dates, vals, d):
    """Linear interpolation of an average onto date d (None outside its range)."""
    import bisect
    if not dates or d < dates[0] or d > dates[-1]:
        return None
    i = bisect.bisect_left(dates, d)
    if dates[i] == d:
        return vals[i]
    d0, d1 = datetime.date.fromisoformat(dates[i - 1]), datetime.date.fromisoformat(dates[i])
    t = (datetime.date.fromisoformat(d) - d0).days / max(1, (d1 - d0).days)
    return round(vals[i - 1] + t * (vals[i] - vals[i - 1]), 2)


def with_dots(avgs, dots, start=None):
    """avgs: [(label, dates, values)] smooth lines; dots: [(label, [(date, value)])] individual polls.
    Returns a table on the union of dates: lines interpolated (so they stay continuous), dots only on
    the dates a poll exists (Datawrapper draws them as symbols with an invisible line)."""
    dd = [dict((d, v) for d, v in pts if v is not None) for _, pts in dots]
    dates = sorted({d for _, ds, _ in avgs for d in ds} | {d for m in dd for d in m})
    if start:
        dates = [d for d in dates if d >= start]
    rows = []
    for d in dates:
        row = [d] + [_interp(ds, vs, d) for _, ds, vs in avgs] + [m.get(d) for m in dd]
        rows.append(row)
    return table(["Fecha"] + [a[0] for a in avgs] + [x[0] for x in dots], rows)


def dot_label(name):
    return f"{name} · encuestas"


def runoff():
    e = load("encuestas_history.json") or {}
    a = e.get("avg") or {}
    polls = [p for p in e.get("polls", []) if not p.get("archived")]
    t = with_dots([("Milei", a.get("dates", []), a.get("milei", [])),
                   ("Kicillof", a.get("dates", []), a.get("kicillof", []))],
                  [(dot_label("Milei"), [(p["date"], p.get("milei")) for p in polls]),
                   (dot_label("Kicillof"), [(p["date"], p.get("kicillof")) for p in polls])])
    return t, a


ESPACIO_LABELS = {"La Libertad Avanza": "LLA", "Peronismo": "Peronismo", "PRO": "PRO",
                  "Izquierda": "Izquierda", "Sin definir": "Indecisos", "Alianza LLA+PRO": "LLA + PRO"}


def espacio(spaces):
    s = load("espacio_history.json") or {"spaces": {}}
    avgs, dots = [], []
    for sp in spaces:
        blk = s["spaces"].get(sp, {})
        av = blk.get("avg", {})
        avgs.append((ESPACIO_LABELS[sp], av.get("dates", []), av.get("values", [])))
        dots.append((dot_label(ESPACIO_LABELS[sp]), [(p["date"], p["value"]) for p in blk.get("polls", [])]))
    return with_dots(avgs, dots), s


def approval(months=12):
    a = load("approval_history.json") or {"avg": {"dates": []}}
    av = a["avg"]
    start = ago(months * 31)
    polls = a.get("polls", [])
    t = with_dots([("Aprueba", av["dates"], av["approve"]), ("Desaprueba", av["dates"], av["disapprove"])],
                  [(dot_label("Aprueba"), [(p["date"], p.get("approve")) for p in polls]),
                   (dot_label("Desaprueba"), [(p["date"], p.get("disapprove")) for p in polls])], start=start)
    return t, a


COMP_ES = {"General evaluation of government": "Evaluación general",
           "Capacity to solve problems": "Capacidad para resolver problemas",
           "Honesty of officials": "Honestidad",
           "Efficiency of public spending": "Eficiencia del gasto",
           "Concern for the general interest": "Preocupación por el interés general"}


def icg():
    d = load("data.json")["icg"]
    t1 = table(["Mes", "ICG"], [(p["date"], p["value"]) for p in d["index"]])
    c = d["components"]
    keys = list(c["series"])
    t2 = table(["Mes"] + [COMP_ES.get(k, k) for k in keys],
               [[m] + [c["series"][k][i] for k in keys] for i, m in enumerate(c["months"])])
    return t1, t2, d


PRES_ES = {"Nestor Kirchner": "N. Kirchner", "Cristina Fernandez de Kirchner I": "CFK I",
           "Cristina Fernandez de Kirchner II": "CFK II", "Mauricio Macri": "Macri",
           "Alberto Fernandez": "A. Fernández", "Javier Milei": "Milei"}


def icg_by_president():
    import openpyxl
    ws = openpyxl.load_workbook(HERE / "Evolution_ICG.xlsx", data_only=True).worksheets[0]
    series, cur = {}, None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            cur = str(r[0]).strip()
            series[cur] = {}
        if cur and r[1] is not None and r[2] not in (None, ""):
            series[cur][int(r[1])] = round(float(r[2]), 2)
    milei = next(k for k in series if "Milei" in k)
    maxm = max(series[milei])
    names = list(series)
    rows = [[m] + [series[p].get(m) for p in names] for m in range(0, maxm + 1)]
    return table(["Mes de gestión"] + [PRES_ES.get(p, p) for p in names], rows), series, maxm


ELECTIONS = [  # year, kind, ICG in the election month, incumbent's vote (see index.html)
    (2005, "Legislativa", 2.192, 39.46), (2007, "Presidencial", 1.678, 45.29),
    (2009, "Legislativa", 1.340, 29.89), (2011, "Presidencial", 2.621, 54.11),
    (2013, "Legislativa", 1.976, 33.05), (2015, "Presidencial", 1.909, 37.08),
    (2017, "Legislativa", 2.828, 41.75), (2019, "Presidencial", 2.009, 40.28),
    (2021, "Legislativa", 1.454, 34.17), (2023, "Presidencial", 1.223, 36.78),
    (2025, "Legislativa", 2.101, 40.66)]


def icg_votes(latest_icg):
    xs = [e[2] * 20 for e in ELECTIONS]
    ys = [e[3] for e in ELECTIONS]
    n = len(xs)
    mx, my = sum(xs) / n, sum(ys) / n
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    b = sxy / sxx
    a = my - b * mx
    r = sxy / (sxx * syy) ** 0.5
    rows = [(e[0], e[1], round(e[2] * 20, 1), e[3]) for e in ELECTIONS]
    lx = round(latest_icg * 20, 1)
    return table(["Año", "Tipo", "Voto implícito (ICG×20)", "Voto del oficialismo"], rows), \
        {"a": a, "b": b, "r": r, "implied": lx, "fitted": a + b * lx}


def polymarket(top=4):
    pm = load("polymarket_history.json") or {"candidates": []}
    cands = [c for c in pm["candidates"] if len(c.get("points", [])) > 1][:top]
    dates = sorted({p[0] for c in cands for p in c["points"]})
    cols = [dict((p[0], p[1]) for p in c["points"]) for c in cands]
    rows = [[d] + [c.get(d) for c in cols] for d in dates]
    return table(["Fecha"] + [c["name"] for c in cands], rows), cands


def governors():
    import openpyxl
    ws = openpyxl.load_workbook(HERE / "Governors_CB.xlsx", data_only=True).worksheets[0]
    rows, month = [], None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is None or r[3] is None:
            continue
        rows.append({"rank": r[0], "gov": r[1], "prov": r[2], "pos": float(r[3]), "neg": float(r[4]),
                     "ns": float(r[5]), "prev": (float(r[6]) if r[6] not in (None, "") else None)})
        month = r[7] or month
    rows.sort(key=lambda x: -x["pos"])
    t = table(["Gobernador", "Positiva", "Negativa", "NS/NC"],
              [(f"{x['gov']} ({x['prov']})", x["pos"], x["neg"], x["ns"]) for x in rows])
    return t, rows, month


# ------------------------------------------------------------------------------ Latin America
MIN_POLLS = 2
COUNTRY_ES = {"Chile": "Chile", "Bolivia": "Bolivia", "Peru": "Perú", "Colombia": "Colombia",
              "Venezuela": "Venezuela", "Mexico": "México", "Ecuador": "Ecuador",
              "El Salvador": "El Salvador", "Panama": "Panamá", "Guatemala": "Guatemala",
              "Honduras": "Honduras", "Brazil": "Brasil"}


def latam():
    L = load("latam_approval.json") or {"countries": {}}
    out = []
    for code, k in L["countries"].items():
        polls = k.get("polls", [])
        av = (k.get("avg") or {}).get("approve") or {}
        dv = (k.get("avg") or {}).get("disapprove") or {}
        if len(polls) < MIN_POLLS or not av.get("values"):
            continue
        last_poll = max(polls, key=lambda p: p["date"])
        out.append({"code": code, "country": COUNTRY_ES.get(k["country"], k["country"]),
                    "president": k["president"], "approve": av["values"][-1],
                    "disapprove": dv["values"][-1] if dv.get("values") else None,
                    "approve_30": at(av["dates"], av["values"], ago(30)),
                    "n": len(polls), "n_30": sum(1 for p in polls if p["date"] >= ago(30)),
                    "last": last_poll, "avg": av, "davg": dv, "polls": polls})
    ar = load("approval_history.json")
    if ar and ar.get("latest"):
        A = ar["avg"]
        out.append({"code": "AR", "country": "Argentina", "president": "Javier Milei",
                    "approve": ar["latest"]["approve"], "disapprove": ar["latest"]["disapprove"],
                    "approve_30": at(A["dates"], A["approve"], ago(30)), "n": len(ar.get("polls", [])),
                    "n_30": sum(1 for p in ar.get("polls", []) if p["date"] >= ago(30)),
                    "last": max(ar["polls"], key=lambda p: p["date"]), "avg": None, "davg": None})
    out.sort(key=lambda r: -r["approve"])
    return out


def latam_ranking_table(rows):
    return table(["Presidente", "Aprobación"],
                 [(f"{r['president']} ({r['country']})", round(r["approve"], 1)) for r in rows])


def latam_country_table(r):
    av, dv = r["avg"], r["davg"] or {}
    polls = r.get("polls", [])
    return with_dots([("Aprueba", av.get("dates", []), av.get("values", [])),
                      ("Desaprueba", dv.get("dates", []), dv.get("values", []))],
                     [(dot_label("Aprueba"), [(p["date"], p.get("approve")) for p in polls]),
                      (dot_label("Desaprueba"), [(p["date"], p.get("disapprove")) for p in polls])])
